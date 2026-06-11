from app.domain_adapter.pdf_extractor import DocumentExtractor


class PPTExtractor(DocumentExtractor):
    def extract(self, path: str) -> str:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RuntimeError("缺少 python-pptx 依赖，请先安装 python-pptx") from exc

        prs = Presentation(path)
        segments: list[str] = []
        for index, slide in enumerate(prs.slides, start=1):
            lines: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    lines.extend(self._clean_lines(shape.text))
                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        for cell in row.cells:
                            lines.extend(self._clean_lines(cell.text))
            notes = getattr(slide, "notes_slide", None)
            if notes and notes.notes_text_frame and notes.notes_text_frame.text:
                lines.extend(self._clean_lines(notes.notes_text_frame.text))
            if lines:
                segments.append(f"[Slide {index}]\n" + "\n".join(lines))
        return "\n\n".join(segments)

    def _clean_lines(self, text: str) -> list[str]:
        return [line.strip() for line in text.splitlines() if line.strip()]
